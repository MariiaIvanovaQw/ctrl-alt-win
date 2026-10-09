"""
Пояснительная записка в Word: те же блоки, что у PDF (scripts/build_docs.py).

Заголовки оформлены встроенными стилями Word («Заголовок 1-3»): по ним
работают область навигации и автоматическое оглавление. Вызывается из
`python -m scripts.build_docs`; нужен python-docx из requirements-dev.txt.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from scripts.markdown_blocks import BOLD, FOOTER, INLINE_CODE, LINK, SUBTITLE, TITLE, Block

BRAND = RGBColor(0x40, 0x2F, 0xFF)  # фиолетовый из брендбука ФСП
HEADING = RGBColor(0x1A, 0x1F, 0x36)
LINK_COLOR = RGBColor(0x1F, 0x3C, 0x88)
MUTED = RGBColor(0x5B, 0x64, 0x75)
FONT, MONO = "Arial", "Consolas"
TEXT_WIDTH_MM = 210 - 2 * 18
PLAIN = re.compile("%s|%s" % (BOLD.pattern, LINK.pattern))


def _shade(element, fill: str) -> None:
    """Заливка ячейки или абзаца (w:shd)."""
    props = element.get_or_add_tcPr() if element.tag == qn("w:tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def _set_font(style, name: str, size: float, bold: bool | None = None, color: RGBColor | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    # кириллица и «восточноазиатский» шрифт тоже Arial, иначе Word подставит Times
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _setup(doc: Document) -> None:
    # шаблон python-docx помечен как Word 2010: без этого Word открывает файл в «режиме совместимости»
    for setting in doc.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(16)
    normal = doc.styles["Normal"]
    _set_font(normal, FONT, 10)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in ((1, 15), (2, 12.5), (3, 11)):
        style = doc.styles["Heading %d" % level]
        _set_font(style, FONT, size, bold=True, color=BRAND if level < 3 else HEADING)
        style.paragraph_format.space_before = Pt(12 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Paragraph"):
        _set_font(doc.styles[name], FONT, 10)
        doc.styles[name].paragraph_format.space_after = Pt(2)


def _footer(doc: Document) -> None:
    """Подвал: название документа и номер страницы (поле PAGE); на титульном листе подвала нет."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    p = section.footer.paragraphs[0]
    p.text = ""
    run = p.add_run(FOOTER + "\t\t")
    run.font.size, run.font.color.rgb = Pt(7.5), MUTED
    for kind in ("begin", "instr", "end"):
        r = p.add_run()
        r.font.size, r.font.color.rgb = Pt(7.5), MUTED
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r._r.append(el)


def _hyperlink(paragraph, url: str, text: str, size: float | None) -> None:
    rel = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                   is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F3C88")
    rpr.append(color)  # порядок в w:rPr по схеме: color, sz, u
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text: str, size: float | None = None, bold: bool = False, italic: bool = False,
               color: RGBColor | None = None) -> None:
    """Строка Markdown → прогоны абзаца: `код`, **полужирный**, [ссылки](адрес)."""

    def run(chunk: str, strong: bool = False, mono: bool = False) -> None:
        if not chunk:
            return
        r = paragraph.add_run(chunk)
        r.bold = (strong or bold) or None
        r.italic = italic or None
        if size:
            r.font.size = Pt(size)
        if mono:
            r.font.name = MONO
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), MONO)
            r.font.color.rgb = LINK_COLOR
        elif color is not None:
            r.font.color.rgb = color

    def plain(chunk: str) -> None:
        last = 0
        for m in PLAIN.finditer(chunk):
            run(chunk[last:m.start()])
            if m.group(1) is not None:
                run(m.group(1), strong=True)
            elif m.group(3).startswith("http"):
                _hyperlink(paragraph, m.group(3), m.group(2), size)
            else:
                run(m.group(2))  # ссылка на файл репозитория — просто текст
            last = m.end()
        run(chunk[last:])

    last = 0
    for m in INLINE_CODE.finditer(text):
        plain(text[last:m.start()])
        run(m.group(1), mono=True)
        last = m.end()
    plain(text[last:])


def _table(doc: Document, rows: list[list[str]]) -> None:
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # ширина столбцов — по длине текста, как в PDF
    lengths = [max(min(len(r[i]), 60) for r in rows) + 4 for i in range(ncols)]
    widths = [max(18.0, TEXT_WIDTH_MM * n / sum(lengths)) for n in lengths]
    widths = [Mm(w * TEXT_WIDTH_MM / sum(widths)) for w in widths]
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.width = widths[c]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, value, size=8.5, bold=r == 0)
            if r == 0:
                _shade(cell._tc, "E8EDF7")
    # первая строка повторяется на каждой странице
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    _shade(p._p, "F5F6F8")  # до интервалов: в w:pPr заливка идёт раньше w:spacing
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name, r.font.size = MONO, Pt(7.5)
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), MONO)
        if i < len(lines) - 1:
            r.add_break(WD_BREAK.LINE)


def _list(doc: Document, items: list[tuple[int, str, str]]) -> None:
    for level, marker, text in items:
        if marker[0].isdigit():
            # номер из исходника сохраняется как есть: Word иначе продолжил бы чужую нумерацию
            p = doc.add_paragraph(style="List Paragraph")
            p.paragraph_format.left_indent = Mm(6 + 6 * level)
            p.paragraph_format.first_line_indent = Mm(-5)
            p.add_run(marker + " ")
        else:
            p = doc.add_paragraph(style=("List Bullet", "List Bullet 2", "List Bullet 3", "List Bullet 3")[level])
        add_inline(p, text)


def add_blocks(doc: Document, blocks: list[Block]) -> None:
    for block in blocks:
        if block.kind == "heading":
            add_inline(doc.add_heading(level=block.level), block.text)
        elif block.kind == "paragraph":
            add_inline(doc.add_paragraph(), block.text)
        elif block.kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Mm(6)
            add_inline(p, block.text, italic=True, color=MUTED)
        elif block.kind == "code":
            _code(doc, block.text)
        elif block.kind == "table":
            _table(doc, block.rows)
        elif block.kind == "list":
            _list(doc, block.rows)


def build_docx(parts: list[tuple[str, list[Block]]], out: Path, date_text: str) -> None:
    doc = Document()
    _setup(doc)
    _footer(doc)
    doc.core_properties.title = "%s – %s" % (TITLE, SUBTITLE.lower())
    doc.core_properties.author = ""

    for _ in range(8):
        doc.add_paragraph()
    for text, size, color in ((TITLE, 22, BRAND), (SUBTITLE, 13, MUTED), (date_text, 11, MUTED)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = size > 20, Pt(size), color
    doc.add_page_break()

    toc = doc.add_paragraph().add_run("Содержание")
    toc.bold, toc.font.size, toc.font.color.rgb = True, Pt(15), BRAND
    for title, _blocks in parts:
        doc.add_paragraph(title).paragraph_format.space_after = Pt(3)
    doc.add_page_break()
    for _title, blocks in parts:
        add_blocks(doc, blocks)
    doc.save(str(out))
